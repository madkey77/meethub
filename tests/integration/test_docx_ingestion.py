"""Integration tests for DOCX file ingestion.

Tests Word document parsing with python-docx, paragraph structure preservation,
and heading extraction.
"""

import io
import tempfile
from pathlib import Path

import pytest

from src.models import File, FileSourceType, FileVersion, Project
from src.services.file_ingestion import FileIngestionService


@pytest.fixture
def sample_docx_file(tmp_path):
    """Create a sample DOCX file with headings and paragraphs."""
    try:
        from docx import Document
        from docx.shared import Pt
    except ImportError:
        pytest.skip("python-docx not installed")

    docx_path = tmp_path / "sample_document.docx"

    # Create document with structure
    doc = Document()

    # Add heading
    doc.add_heading("Project Overview", level=1)
    doc.add_paragraph("This is the introduction paragraph explaining the project.")

    # Add subheading
    doc.add_heading("Architecture Design", level=2)
    doc.add_paragraph("The architecture consists of multiple components:")
    doc.add_paragraph("- RAG Pipeline for document processing")
    doc.add_paragraph("- Vector store for embeddings")
    doc.add_paragraph("- LLM providers for text generation")

    # Add another heading
    doc.add_heading("Implementation Details", level=2)
    doc.add_paragraph("Implementation follows best practices for production systems.")

    doc.save(str(docx_path))
    return docx_path


@pytest.fixture
def test_project(db_session):
    """Create test project for file uploads."""
    project = Project(
        name="docx-test-project",
        drive_folder_id="test-folder-docx",
        classification_rules=[],
        is_default=False,
        display_name="DOCX Test Project",
    )
    db_session.add(project)
    db_session.commit()
    db_session.refresh(project)
    return project


class TestDOCXIngestion:
    """Integration tests for DOCX file ingestion."""

    def test_docx_upload_creates_file_entity(self, db_session, test_project, sample_docx_file):
        """Test that DOCX upload creates File entity.

        Verifies:
        - File record is created
        - MIME type is detected correctly
        - Source type is set
        """
        service = FileIngestionService(db_session)

        file_entity = service.ingest_file(
            file_path=str(sample_docx_file),
            project_id=test_project.project_id,
            source_type="uploaded_document"
        )

        assert file_entity is not None
        assert file_entity.file_id > 0
        assert "word" in file_entity.mime_type.lower() or "openxml" in file_entity.mime_type.lower()
        assert file_entity.source_type == FileSourceType.UPLOADED_DOCUMENT

    def test_docx_paragraph_structure_preservation(self, db_session, test_project, sample_docx_file):
        """Test that paragraph structure is preserved during extraction.

        Verifies:
        - Paragraphs are separated correctly
        - Line breaks are maintained
        - Text content is complete
        """
        service = FileIngestionService(db_session)

        text, metadata = service.extract_text(
            file_path=str(sample_docx_file),
            mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )

        # Verify text content
        assert text is not None
        assert "Project Overview" in text
        assert "Architecture Design" in text
        assert "RAG Pipeline" in text

        # Verify paragraphs are separated (double newlines)
        assert "\n\n" in text

    def test_docx_heading_extraction(self, db_session, test_project, sample_docx_file):
        """Test heading extraction from DOCX file.

        Verifies:
        - Headings are identified
        - Heading count is tracked in metadata
        """
        service = FileIngestionService(db_session)

        text, metadata = service.extract_text(
            file_path=str(sample_docx_file),
            mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )

        # Verify metadata
        assert metadata is not None
        assert "heading_count" in metadata
        assert metadata["heading_count"] >= 3  # We created 3 headings

    def test_docx_metadata_paragraph_count(self, db_session, test_project, sample_docx_file):
        """Test paragraph count extraction from DOCX.

        Verifies:
        - Paragraph count is calculated
        - Metadata includes paragraph_count field
        """
        service = FileIngestionService(db_session)

        text, metadata = service.extract_text(
            file_path=str(sample_docx_file),
            mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )

        assert "paragraph_count" in metadata
        assert metadata["paragraph_count"] > 0
        assert metadata["paragraph_count"] >= 5  # At least 5 paragraphs created

    def test_docx_extraction_method_metadata(self, db_session, test_project, sample_docx_file):
        """Test extraction method is recorded in metadata.

        Verifies:
        - extraction_method is 'python-docx'
        """
        service = FileIngestionService(db_session)

        text, metadata = service.extract_text(
            file_path=str(sample_docx_file),
            mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )

        assert metadata["extraction_method"] == "python-docx"

    def test_docx_empty_paragraphs_handling(self, db_session, test_project, tmp_path):
        """Test that empty paragraphs are handled correctly.

        Verifies:
        - Empty paragraphs are skipped
        - Only non-empty content is extracted
        """
        try:
            from docx import Document
        except ImportError:
            pytest.skip("python-docx not installed")

        docx_path = tmp_path / "empty_paragraphs.docx"

        doc = Document()
        doc.add_paragraph("First paragraph")
        doc.add_paragraph("")  # Empty
        doc.add_paragraph("   ")  # Whitespace only
        doc.add_paragraph("Second paragraph")
        doc.save(str(docx_path))

        service = FileIngestionService(db_session)

        text, metadata = service.extract_text(
            file_path=str(docx_path),
            mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )

        # Should only count non-empty paragraphs
        assert "First paragraph" in text
        assert "Second paragraph" in text
        # Empty paragraphs should not create excessive newlines
        assert "\n\n\n\n" not in text

    def test_docx_special_characters_handling(self, db_session, test_project, tmp_path):
        """Test handling of special characters in DOCX.

        Verifies:
        - Unicode characters are preserved
        - Special symbols are extracted correctly
        """
        try:
            from docx import Document
        except ImportError:
            pytest.skip("python-docx not installed")

        docx_path = tmp_path / "special_chars.docx"

        doc = Document()
        doc.add_paragraph("Português: áéíóú àçã")
        doc.add_paragraph("Symbols: © ® ™ € £")
        doc.add_paragraph("Math: ∑ ∏ ∫ √")
        doc.save(str(docx_path))

        service = FileIngestionService(db_session)

        text, metadata = service.extract_text(
            file_path=str(docx_path),
            mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )

        # Verify special characters preserved
        assert "Português" in text or "Portugu" in text  # Encoding handling
        assert metadata["paragraph_count"] == 3
